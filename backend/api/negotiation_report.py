"""GET /api/rfqs/{rfqnum}/negotiation-report (JSON) and
.../negotiation-report.xlsx (download) -- a per-vendor negotiation-prep
summary combining backend/api/comparison.py's BOQ flags (unquoted,
zero-priced, arithmetic error, outlier) and backend/api/rounds.py's
movement flags (price increase, steeper-than-peer discount).

Deliberately NOT LLM-generated -- no narrative, no "recommended award",
no reasoning prose. Vendors are ranked by contract_total (ascending) and
each gets a capped, deduped list of their most significant issues; the
analyst draws their own conclusions from the ranked facts. This mirrors
full-feature-buildout's backend/reporting/digest.py pattern (compact
per-bidder digest: ranked totals, flag counts, top-N flags deduped by
category before backfilling) without its llm_report.py's Azure OpenAI
call -- that's a deliberate, discussed scope cut for this pass, not an
oversight.

Pure read -- calls build_comparison()/build_round_trend() directly
in-process (no HTTP round trip to our own API), no write-grant
dependency.
"""

from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font

from backend.api.comparison import build_comparison
from backend.api.rounds import build_round_trend

router = APIRouter()

TOP_ISSUES_PER_VENDOR = 8

# Lower number = higher priority when picking which issues make the cut.
_PRIORITY = {
    "round_critical": 0,
    "arithmetic_error": 1,
    "round_warning": 2,
    "unquoted": 3,
    "outlier": 4,
    "zero_price": 5,
}


@dataclass
class TopIssue:
    source: str  # "boq" | "round"
    severity: str
    rfqlinenum: float
    description: str | None
    detail: str


@dataclass
class NegotiationVendorSummary:
    vendor: str
    name: str | None
    rank: int
    contract_total: float
    pct_above_lowest: float
    unquoted_count: int
    arithmetic_error_count: int
    outlier_count: int
    zero_price_count: int
    round_critical_count: int
    round_warning_count: int
    top_issues: list[TopIssue]


def _iso(dt) -> str:
    return dt.isoformat() if hasattr(dt, "isoformat") else str(dt)


def _boq_issue_detail(category: str, qty: float | None, unit_cost, line_cost) -> str:
    if category == "unquoted":
        return "No price submitted for this line."
    if category == "zero_price":
        return "Quoted AED 0 for this line -- likely missing, not a genuine free item."
    if category == "arithmetic_error":
        if qty is not None and unit_cost is not None and line_cost is not None:
            return (
                f"Line total AED {line_cost:,.0f} doesn't match qty x unit rate "
                f"(AED {unit_cost:,.0f} x {qty:,.0f})."
            )
        return "Line total doesn't match qty x unit rate."
    if category == "outlier":
        return f"AED {unit_cost:,.0f}/unit is unusually far from other bidders on this line."
    return ""


def _collect_boq_issues(comparison: dict, vendor: str) -> list[dict]:
    # comparison["lines"] is capped at comparison.py's LINES_CAP (500) --
    # on a BOQ larger than that, issues past the cap are invisible here
    # too, the same limitation the comparison view itself already has.
    issues = []
    for line in comparison["lines"]:
        p = line["prices"].get(vendor)
        if not p:
            continue
        rfqlinenum = line["rfqlinenum"]
        description = line["description"]
        qty = line["qty"]
        unit_cost = p["unit_cost"]
        line_cost = p["line_cost"]

        flagged = [
            (category, "critical" if category == "arithmetic_error" else "warning")
            for category, is_flagged in (
                ("arithmetic_error", p["arithmetic_error"]),
                ("unquoted", p["unquoted"]),
                ("outlier", p["outlier"]),
                ("zero_price", p["zero_price"]),
            )
            if is_flagged
        ]
        for category, severity in flagged:
            issues.append(
                dict(
                    source="boq",
                    category=category,
                    priority=_PRIORITY[category],
                    severity=severity,
                    rfqlinenum=rfqlinenum,
                    description=description,
                    detail=_boq_issue_detail(category, qty, unit_cost, line_cost),
                )
            )
    return issues


def _collect_round_issues(round_trend: dict, vendor: str) -> list[dict]:
    issues = []
    for f in round_trend["flags"]:
        if f["vendor"] != vendor:
            continue
        category = "round_critical" if f["severity"] == "critical" else "round_warning"
        issues.append(
            dict(
                source="round",
                category=category,
                priority=_PRIORITY[category],
                severity=f["severity"],
                rfqlinenum=f["rfqlinenum"],
                description=f["description"],
                detail=f["detail"],
            )
        )
    return issues


def _pick_top_issues(issues: list[dict], cap: int) -> list[dict]:
    """Category-first dedup, then backfill -- so a lone arithmetic error
    isn't crowded out by five outliers, mirroring full-feature-buildout's
    digest.py::_top_flags."""
    picked_idx: list[int] = []
    seen_categories: set[str] = set()
    for i, issue in enumerate(issues):
        if issue["category"] not in seen_categories:
            picked_idx.append(i)
            seen_categories.add(issue["category"])
        if len(picked_idx) >= cap:
            return [issues[i] for i in picked_idx]
    for i, issue in enumerate(issues):
        if i in picked_idx:
            continue
        picked_idx.append(i)
        if len(picked_idx) >= cap:
            break
    return [issues[i] for i in picked_idx]


def build_negotiation_report(rfqnum: str) -> dict:
    comparison = build_comparison(rfqnum)
    round_trend = build_round_trend(rfqnum)

    has_round_data = len(round_trend["rounds_present"]) >= 2

    round_critical_counts: dict[str, int] = {}
    round_warning_counts: dict[str, int] = {}
    for f in round_trend["flags"]:
        counts = round_critical_counts if f["severity"] == "critical" else round_warning_counts
        counts[f["vendor"]] = counts.get(f["vendor"], 0) + 1

    ranked = sorted(comparison["vendors"], key=lambda v: v["contract_total"])
    lowest_total = ranked[0]["contract_total"] if ranked else None

    vendors = []
    for rank, v in enumerate(ranked, start=1):
        vendor = v["vendor"]
        pct_above_lowest = (
            round((v["contract_total"] - lowest_total) / lowest_total * 100, 2)
            if lowest_total
            else 0.0
        )

        raw_issues = _collect_boq_issues(comparison, vendor) + _collect_round_issues(round_trend, vendor)
        raw_issues.sort(key=lambda i: (i["priority"], i["rfqlinenum"]))
        top = _pick_top_issues(raw_issues, TOP_ISSUES_PER_VENDOR)

        vendors.append(
            NegotiationVendorSummary(
                vendor=vendor,
                name=v["name"],
                rank=rank,
                contract_total=v["contract_total"],
                pct_above_lowest=pct_above_lowest,
                unquoted_count=v["unquoted_count"],
                arithmetic_error_count=v["arithmetic_error_count"],
                outlier_count=v["outlier_count"],
                zero_price_count=v["zero_price_count"],
                round_critical_count=round_critical_counts.get(vendor, 0),
                round_warning_count=round_warning_counts.get(vendor, 0),
                top_issues=[
                    TopIssue(
                        source=i["source"],
                        severity=i["severity"],
                        rfqlinenum=i["rfqlinenum"],
                        description=i["description"],
                        detail=i["detail"],
                    )
                    for i in top
                ],
            )
        )

    return {
        "rfqnum": rfqnum,
        "generated_at": _iso(datetime.now(timezone.utc)),
        "has_round_data": has_round_data,
        "vendors": [asdict(v) for v in vendors],
        "not_submitted": comparison["not_submitted"],
    }


def _build_excel(report: dict) -> BytesIO:
    wb = openpyxl.Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(
        [
            "Vendor", "Name", "Rank", "Contract Total", "% Above Lowest",
            "Unquoted", "Arithmetic Errors", "Outliers", "Zero-Priced",
            "Round Critical", "Round Warning",
        ]
    )
    for cell in ws_summary[1]:
        cell.font = Font(bold=True)
    for v in report["vendors"]:
        ws_summary.append(
            [
                v["vendor"], v["name"] or "", v["rank"], v["contract_total"], v["pct_above_lowest"],
                v["unquoted_count"], v["arithmetic_error_count"], v["outlier_count"], v["zero_price_count"],
                v["round_critical_count"], v["round_warning_count"],
            ]
        )
    for row in ws_summary.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = "#,##0"

    ws_issues = wb.create_sheet("Issues")
    ws_issues.append(["Vendor", "Name", "Source", "Severity", "Line", "Description", "Detail"])
    for cell in ws_issues[1]:
        cell.font = Font(bold=True)
    for v in report["vendors"]:
        for issue in v["top_issues"]:
            ws_issues.append(
                [
                    v["vendor"], v["name"] or "", issue["source"], issue["severity"], issue["rfqlinenum"],
                    issue["description"] or "", issue["detail"],
                ]
            )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.get("/rfqs/{rfqnum}/negotiation-report")
async def get_negotiation_report(rfqnum: str):
    return build_negotiation_report(rfqnum)


@router.get("/rfqs/{rfqnum}/negotiation-report.xlsx")
async def get_negotiation_report_excel(rfqnum: str):
    report = build_negotiation_report(rfqnum)
    buf = _build_excel(report)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{rfqnum}-negotiation-report.xlsx"'},
    )
