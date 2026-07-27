"""Extract structured BOQ data from Excel files."""

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl


@dataclass(frozen=True)
class BOQItem:
    item_no: str
    description: str
    unit: str | None
    qty: float | None
    cif_unit_rate: float | None
    cif_total: float | None
    erection_unit_rate: float | None
    erection_total: float | None
    total: float | None
    is_section_header: bool = False
    raw_cif: str | None = None  # original cell value (for "Included", "N/A" etc.)
    raw_erection: str | None = None
    is_corrected: bool = False  # True if a user overrode the OCR/extracted value
    is_missing: bool = False  # True if no source data exists at all (data-gap skeleton item)


@dataclass(frozen=True)
class BOQSheet:
    name: str
    items: tuple[BOQItem, ...]


@dataclass(frozen=True)
class BOQLot:
    lot_name: str
    lot_number: int
    sheets: tuple[BOQSheet, ...]
    total_cif: float | None = None
    total_erection: float | None = None
    total_price: float | None = None


@dataclass(frozen=True)
class BOQExtraction:
    tender_no: str
    bidder: str
    round_name: str
    lots: tuple[BOQLot, ...]
    total_contract_price: float | None = None
    # Set when no genuine per-item BOQ source exists for this bidder/round (e.g.
    # only a lot-summary document was submitted) — lots/items are then a
    # skeleton cloned from a reference bidder's item structure, all pricing
    # fields None and is_missing=True, awaiting manual reviewer entry.
    data_gap: str | None = None


INCLUDED_MARKERS = {"included", "inc", "incl", "included above", "included in"}
NOT_APPLICABLE = {"", "-", "n/a", "na", "nil", "no boq item"}


def _safe_float(val) -> float | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        cleaned = val.strip().replace(",", "")
        lower = cleaned.lower()
        if lower in NOT_APPLICABLE:
            return None
        # "Included" means price is bundled in another item — treat as explicit 0
        if lower in INCLUDED_MARKERS or lower.startswith("included"):
            return 0.0
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def _safe_str(val) -> str | None:
    if val is None:
        return None
    return str(val).strip() if str(val).strip() else None


def _effective_col_count(ws) -> int:
    """openpyxl's ws.max_column can be inflated by stray content far below
    the header -- e.g. corrupted #REF! formula artifacts sitting on
    subtotal rows -- and column count drives which fields (unit/qty/rate/
    erection/total) get read for every row in the sheet, so an inflated
    count silently misreads real data as blank (seen on a bidder whose
    Lot 2 "Items 9" sheet reported max_column=10 while every genuine row
    only ever populates 6, because a subtotal row's leftover #REF! cells
    happened to land in columns 8-10).

    The real column layout is defined by the header block itself: the row
    with "Item" in column 1, plus the two rows below it (sub-headers like
    "Unit Rate (b)" / "Total (a*b)", and — for the full 9-column layout —
    "A + B" one row further down under the Total column).
    """
    header_row = None
    for r in range(1, min(15, ws.max_row) + 1):
        if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).strip() == "Item":
            header_row = r
            break
    if header_row is None:
        return ws.max_column

    max_col = 0
    for r in range(header_row, header_row + 3):
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=r, column=c).value is not None:
                max_col = max(max_col, c)
    return max_col


def _is_section_header(row_vals: list) -> bool:
    """A section header has an item number and description but no pricing."""
    has_item = row_vals[0] is not None
    has_desc = row_vals[1] is not None
    has_no_pricing = all(v is None for v in row_vals[2:])
    return has_item and has_desc and has_no_pricing


def _find_included_parents_in_sheet(ws, col_count: int) -> set[str]:
    """Scan for 'Included in BOQ item X' patterns to find parent items.

    When a sub-item says "Included in BOQ item 3.3", all sibling sub-items
    under 3.3 are implicitly bundled into the parent's price.
    """
    parents: set[str] = set()
    # Check pricing columns for "Included in" text
    price_cols = [4, 5, 8] if col_count >= 9 else [4, 5]
    for row in ws.iter_rows(min_row=1, values_only=True):
        for col_idx in price_cols:
            if col_idx < len(row):
                val = row[col_idx]
                if val and isinstance(val, str) and "included in" in val.lower():
                    m = re.search(r"item\s+(\d[\d.]*)", val, re.IGNORECASE)
                    if m:
                        parents.add(m.group(1))
    return parents


def _is_sub_item_of(item_no: str, parents: set[str]) -> bool:
    """Check if item_no is a sub-item of any included parent."""
    for parent in parents:
        if item_no.startswith(parent + ".") and item_no != parent:
            return True
    return False


def _parse_detail_sheet(ws, col_count: int) -> list[BOQItem]:
    """Parse a detail sheet (Items 1-3, Item 4, Items 5-8, or Item 9)."""
    items = []
    header_row = None
    included_parents = _find_included_parents_in_sheet(ws, col_count)

    for row_idx, row in enumerate(ws.iter_rows(min_row=1, values_only=True), 1):
        # Find the header row with "Item" / "Description"
        if row[0] and str(row[0]).strip() == "Item":
            header_row = row_idx
            continue

        # Skip rows before and just after header (sub-headers, disclaimers)
        if header_row is None:
            continue
        if row_idx <= header_row + 2:
            continue

        # Skip fully empty rows
        if all(v is None for v in row):
            continue

        item_no = _safe_str(row[0])
        description = _safe_str(row[1])

        # Skip disclaimer/text-only rows
        if item_no is None and description and len(str(description)) > 200:
            continue

        # Skip sub-total/total rows
        if description and any(
            kw in str(description).lower()
            for kw in ["sub total", "subtotal", "total price", "total contract",
                        "grand total", "total lot", "u.a.e. dirhams"]
        ):
            continue

        # Skip total rows in item_no field (e.g., "Item-1 Sub Total", recap rows)
        if item_no:
            item_lower = item_no.lower().strip()
            if any(kw in item_lower for kw in ("sub total", "grand total", "total")):
                continue
            if re.match(r"^item\s*-?\s*\d+$", item_lower):
                continue

        if item_no is None and description is None:
            continue

        # Determine column layout based on sheet column count
        if col_count >= 9:
            # Full layout: Item, Desc, Unit, Qty, CIF Rate, CIF Total, Erect Rate, Erect Total, Total
            unit = _safe_str(row[2]) if len(row) > 2 else None
            qty = _safe_float(row[3]) if len(row) > 3 else None
            cif_rate = _safe_float(row[4]) if len(row) > 4 else None
            cif_total = _safe_float(row[5]) if len(row) > 5 else None
            erect_rate = _safe_float(row[6]) if len(row) > 6 else None
            erect_total = _safe_float(row[7]) if len(row) > 7 else None
            total = _safe_float(row[8]) if len(row) > 8 else None
            raw_cif = _safe_str(row[4]) if len(row) > 4 else None
            raw_erect = _safe_str(row[6]) if len(row) > 6 else None
        elif col_count >= 6:
            # Spare parts: Item, Desc, Unit, Qty, CIF Rate, CIF Total
            unit = _safe_str(row[2]) if len(row) > 2 else None
            qty = _safe_float(row[3]) if len(row) > 3 else None
            cif_rate = _safe_float(row[4]) if len(row) > 4 else None
            cif_total = _safe_float(row[5]) if len(row) > 5 else None
            erect_rate = None
            erect_total = None
            total = cif_total
            raw_cif = _safe_str(row[4]) if len(row) > 4 else None
            raw_erect = None
        else:
            continue

        is_header = (
            item_no is not None
            and description is not None
            and unit is None
            and qty is None
            and cif_rate is None
            and cif_total is None
        )

        # Sub-items under an "Included in BOQ item X" parent with no pricing
        # are implicitly bundled — treat as 0.0 (not unquoted)
        has_no_pricing = all(v is None for v in (cif_rate, cif_total, erect_rate, erect_total, total))
        bundled = (
            has_no_pricing
            and qty is not None
            and item_no is not None
            and _is_sub_item_of(item_no, included_parents)
        )

        items.append(
            BOQItem(
                item_no=item_no or "",
                description=description or "",
                unit=unit,
                qty=qty,
                cif_unit_rate=0.0 if bundled else cif_rate,
                cif_total=0.0 if bundled else cif_total,
                erection_unit_rate=0.0 if bundled else erect_rate,
                erection_total=0.0 if bundled else erect_total,
                total=0.0 if bundled else total,
                is_section_header=is_header,
                raw_cif="Bundled in parent" if bundled else raw_cif,
                raw_erection="Bundled in parent" if bundled else raw_erect,
            )
        )

    return items


def _parse_summary_sheet(ws) -> dict:
    """Parse the Final Summary sheet to get lot-level totals."""
    summary = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if row[0] and str(row[0]).strip().lower() == "total price lot wise(u.a.e. dirhams)":
            continue
        if row[1] and "total contract price" in str(row[1]).lower():
            # Find the first non-None numeric value
            for val in row[2:]:
                f = _safe_float(val)
                if f is not None:
                    summary["total_contract_price"] = f
                    break
    return summary


def _detect_lot_info(ws) -> tuple[str, int]:
    """Detect lot name and number from sheet content."""
    for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
        text = " ".join(str(v) for v in row if v is not None).upper()
        # Skip the tender title line that contains all lot names
        if "REPLACEMENT OF" in text or "TENDER NO" in text:
            continue
        # Look for explicit LOT designation lines
        if "LOT 1:" in text or "LOT 1 :" in text:
            return "Lot 1: Shobaisi PRY (SHBPRY)", 1
        if "LOT-2:" in text or "LOT 2:" in text or "LOT-2 :" in text or "LOT 2 :" in text:
            return "Lot 2: Defence Residence PRY (DRPRY)", 2
        if "LOT 3:" in text or "LOT 3 :" in text:
            return "Lot 3: Samha PRY (SMHPRY)", 3
    # Fallback: check filename-style patterns in all text
    all_text = ""
    for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
        all_text += " ".join(str(v) for v in row if v is not None).upper() + " "
    if "LOT 1" in all_text and "SHBPRY" in all_text and "DRPRY" not in all_text.split("LOT 1")[0]:
        return "Lot 1: Shobaisi PRY (SHBPRY)", 1
    return "Unknown Lot", 0


def _detect_lot_from_filename(filename: str) -> tuple[str, int]:
    """Detect lot from filename as a fallback."""
    upper = filename.upper()
    if "LOT 1" in upper or "LOT-1" in upper or "SHBPRY" in upper:
        return "Lot 1: Shobaisi PRY (SHBPRY)", 1
    if "LOT 2" in upper or "LOT-2" in upper or "DRPRY" in upper:
        return "Lot 2: Defence Residence PRY (DRPRY)", 2
    if "LOT 3" in upper or "LOT-3" in upper or "SMHPRY" in upper:
        return "Lot 3: Samha PRY (SMHPRY)", 3
    return "Unknown Lot", 0


def parse_lot_file(filepath: str | Path) -> BOQLot:
    """Parse a single lot Excel file into structured BOQ data."""
    filepath = Path(filepath)
    wb = openpyxl.load_workbook(str(filepath), data_only=True)
    sheets = []
    lot_name = ""
    lot_number = 0
    total_cif = None
    total_erection = None
    total_price = None

    # Detect lot from filename first (most reliable)
    lot_name, lot_number = _detect_lot_from_filename(filepath.name)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        if "final summary" in sheet_name.lower():
            if lot_number == 0:
                lot_name, lot_number = _detect_lot_info(ws)
            # Extract totals from summary — look for the total row
            for row in ws.iter_rows(min_row=1, values_only=True):
                desc = _safe_str(row[1]) if len(row) > 1 else None
                if desc and "total" in desc.lower():
                    cif = _safe_float(row[2]) if len(row) > 2 else None
                    erect = _safe_float(row[3]) if len(row) > 3 else None
                    price = _safe_float(row[4]) if len(row) > 4 else None
                    # Take the last matching "total" row (which is the grand total)
                    if cif is not None or erect is not None or price is not None:
                        total_cif = cif
                        total_erection = erect
                        total_price = price
            continue

        items = _parse_detail_sheet(ws, _effective_col_count(ws))
        if items:
            sheets.append(BOQSheet(name=sheet_name, items=tuple(items)))

    return BOQLot(
        lot_name=lot_name,
        lot_number=lot_number,
        sheets=tuple(sheets),
        total_cif=total_cif,
        total_erection=total_erection,
        total_price=total_price,
    )


def parse_summary_file(filepath: str | Path) -> dict:
    """Parse the summary-of-lots file for contract-level totals."""
    wb = openpyxl.load_workbook(str(filepath), data_only=True)
    ws = wb[wb.sheetnames[0]]
    result = {"lots": {}, "total_contract_price": None}

    for row in ws.iter_rows(min_row=1, values_only=True):
        if row[1] and "total contract price" in str(row[1]).lower():
            for val in row[2:]:
                f = _safe_float(val)
                if f is not None:
                    result["total_contract_price"] = f
                    break
        if row[1] and "total price lot wise" in str(row[1]).lower():
            # Lot 1 totals in cols 2-4, Lot 2 in 5-7, Lot 3 in 8-10
            result["lots"] = {
                1: _safe_float(row[4]) if len(row) > 4 else None,
                2: _safe_float(row[7]) if len(row) > 7 else None,
                3: _safe_float(row[10]) if len(row) > 10 else None,
            }

    return result


ROUND_DIR_NAMES = ("original", "round1", "round2", "round3")


def _parse_round_dir(round_dir: Path, bidder_name: str) -> BOQExtraction | None:
    """Parse a single round directory's Excel files into a BOQExtraction."""
    excel_dir = round_dir / "excel"
    if not excel_dir.exists():
        return None

    round_name = round_dir.name
    lots = []
    total_contract_price = None

    for xlsx_file in sorted(excel_dir.glob("*.xlsx")):
        fname = xlsx_file.name.lower()
        if "summary" in fname:
            summary = parse_summary_file(xlsx_file)
            total_contract_price = summary.get("total_contract_price")
        elif "lot" in fname or "shbpry" in fname or "drpry" in fname or "smhpry" in fname:
            lot = parse_lot_file(xlsx_file)
            lots.append(lot)

    if not lots:
        return None

    lots.sort(key=lambda l: l.lot_number)

    # Some bidders don't submit a separate summary-of-all-lots file
    # (e.g. only 3 per-lot files, no combined summary) — fall back to
    # summing each lot's own total_price rather than leaving the
    # contract total blank when we already have the real numbers.
    if total_contract_price is None:
        lot_totals = [l.total_price for l in lots if l.total_price is not None]
        if lot_totals:
            total_contract_price = sum(lot_totals)

    return BOQExtraction(
        tender_no="D-111808",
        bidder=bidder_name,
        round_name=round_name,
        lots=tuple(lots),
        total_contract_price=total_contract_price,
    )


def parse_bidder_folder(bidder_path: str | Path, bidder_name: str) -> list[BOQExtraction]:
    """Parse the original-round submission for a bidder (Excel fallback path).

    Original-bidding skeleton: negotiation rounds are intentionally not scanned.
    """
    bidder_path = Path(bidder_path)
    round_dirs = sorted(
        [d for d in bidder_path.iterdir() if d.is_dir() and d.name == "original"],
        key=lambda d: d.name,
    )
    extractions = [_parse_round_dir(d, bidder_name) for d in round_dirs]
    return [e for e in extractions if e is not None]


def parse_bidder_folder_all_rounds(bidder_path: str | Path, bidder_name: str) -> list[BOQExtraction]:
    """Parse every negotiation round (original, round1, round2, round3) found
    for a bidder, via Excel — used for round-over-round price movement
    tracking. Excel is the reliable source across rounds (openpyxl parses it
    cleanly; the PDF path is only wired up for the original round). Rounds
    with no usable Excel data for this bidder are simply omitted, not errored.
    """
    bidder_path = Path(bidder_path)
    round_dirs = sorted(
        [d for d in bidder_path.iterdir() if d.is_dir() and d.name in ROUND_DIR_NAMES],
        key=lambda d: ROUND_DIR_NAMES.index(d.name),
    )
    extractions = [_parse_round_dir(d, bidder_name) for d in round_dirs]
    return [e for e in extractions if e is not None]
