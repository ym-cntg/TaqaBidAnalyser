"""Plain downloadable comparison report — feature 17: "the consolidated
comparison as an Excel file, the artefact the analyst produces by hand
today." Unlike the LLM recommendation report (`excel_export.py`), this has
no AI narrative and includes every detected flag, not a curated top-N —
a raw audit export should show everything, not a summary.
"""

from io import BytesIO

import openpyxl

from ..analysis.comparator import ComparisonResult
from .digest import SEVERITY_RANK, build_report_digest
from .excel_export import (
    BEST_PRICE_FILL,
    DATA_FONT,
    GAP_FILL,
    NUMBER_FORMAT,
    SEVERITY_FILL,
    THIN_BORDER,
    WRAP,
    header_row,
)


def _write_lot_sheet(wb, lot, data_gap_bidders: set[str]) -> None:
    bidders = sorted(lot.bidder_totals.keys(), key=lambda b: b in data_gap_bidders)
    name = f"Lot {lot.lot_number}"[:31]
    ws = wb.create_sheet(name)
    headers = ["Item No", "Description", "Unit", "Qty"] + bidders
    widths = [12, 45, 8, 8] + [18] * len(bidders)
    header_row(ws, 1, headers, widths)

    row = 2
    for item in lot.items:
        totals = {b: item.bidder_prices.get(b, {}).get("total") for b in bidders}
        eligible = [v for b, v in totals.items() if b not in data_gap_bidders and v is not None]
        best = min(eligible) if eligible else None

        cells = [
            ws.cell(row=row, column=1, value=item.item_no),
            ws.cell(row=row, column=2, value=item.description),
            ws.cell(row=row, column=3, value=item.unit),
            ws.cell(row=row, column=4, value=item.qty),
        ]
        for col, bidder in enumerate(bidders, start=5):
            value = totals[bidder]
            cell = ws.cell(row=row, column=col, value=value if value is not None else "—")
            if value is not None:
                cell.number_format = NUMBER_FORMAT
            if bidder in data_gap_bidders:
                cell.fill = GAP_FILL
            elif best is not None and value == best:
                cell.fill = BEST_PRICE_FILL
            cells.append(cell)

        for cell in cells:
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = WRAP
        row += 1

    # Lot totals row
    total_row = row
    ws.cell(row=total_row, column=2, value="LOT TOTAL").font = DATA_FONT
    for col, bidder in enumerate(bidders, start=5):
        total = lot.bidder_totals.get(bidder, {}).get("total")
        cell = ws.cell(row=total_row, column=col, value=total if total is not None else "—")
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        if total is not None:
            cell.number_format = NUMBER_FORMAT


def _write_summary_sheet(wb, result: ComparisonResult, data_gap_bidders: dict[str, str]) -> None:
    ws = wb.create_sheet("Bid Summary", 0)
    header_row(ws, 1, ["Bidder", "Contract Total", "Rank", "Status"], [24, 20, 10, 30])

    eligible = sorted(
        ((b, t) for b, t in result.bidder_grand_totals.items() if t is not None and b not in data_gap_bidders),
        key=lambda bt: bt[1],
    )
    best_bidder = eligible[0][0] if eligible else None
    ranked = [(b, i + 1) for i, (b, _) in enumerate(eligible)] + [(b, None) for b in data_gap_bidders]

    row = 2
    for bidder, rank in ranked:
        total = result.bidder_grand_totals.get(bidder)
        status = data_gap_bidders.get(bidder, "")
        cells = [
            ws.cell(row=row, column=1, value=bidder),
            ws.cell(row=row, column=2, value=total if total is not None else "—"),
            ws.cell(row=row, column=3, value=rank if rank is not None else "—"),
            ws.cell(row=row, column=4, value=status),
        ]
        for cell in cells:
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = WRAP
            if bidder == best_bidder:
                cell.fill = BEST_PRICE_FILL
            elif bidder in data_gap_bidders:
                cell.fill = GAP_FILL
        cells[1].number_format = NUMBER_FORMAT
        row += 1


def _write_all_flags_sheet(wb, result: ComparisonResult) -> None:
    ws = wb.create_sheet("Flags")
    header_row(
        ws, 1,
        ["Severity", "Category", "Bidder", "Lot", "Item No", "Description", "Detail"],
        [10, 16, 20, 16, 10, 40, 60],
    )
    flags = sorted(
        result.flags,
        key=lambda f: (SEVERITY_RANK.get(f.severity, 99), f.bidder, f.category),
    )
    row = 2
    for f in flags:
        values = [f.severity, f.category, f.bidder, f.lot, f.item_no, f.description, f.detail]
        fill = SEVERITY_FILL.get(f.severity)
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = DATA_FONT
            cell.alignment = WRAP
            cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
        row += 1


def build_comparison_excel(result: ComparisonResult) -> BytesIO:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    digest = build_report_digest(result)
    data_gap_bidders = digest["data_gap_bidders"]

    _write_summary_sheet(wb, result, data_gap_bidders)
    for lot in result.lots:
        _write_lot_sheet(wb, lot, set(data_gap_bidders))
    _write_all_flags_sheet(wb, result)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
