"""Excel export of a recommendation report.

Styling constants below are copied by value (not imported) from the frozen
`src/reporting/excel_report.py` reference implementation — that module is
tightly coupled to an incompatible data model (`TenderData`, `Finding`) and
can't be imported directly, but its color/style choices are worth reusing
for visual consistency with the original design intent.
"""

from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..analysis.comparator import ComparisonResult
from .digest import build_report_digest

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
SUBHEADER_FONT = Font(name="Calibri", bold=True, size=10)
DATA_FONT = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
NUMBER_FORMAT = "#,##0"
BEST_PRICE_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
GAP_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
SEVERITY_FILL = {
    "critical": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "warning": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "info": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
}
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def header_row(ws, row: int, headers: list[str], widths: list[int]) -> None:
    for col, (text, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_recommendation_sheet(wb, result: ComparisonResult, report: dict) -> None:
    ws = wb.create_sheet("Recommendation")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100

    rec = report.get("recommendation", {})
    rows = [
        ("Tender", result.tender_no),
        ("Round", result.round_name),
        ("", ""),
        ("Summary", report.get("executive_summary", "")),
        ("", ""),
        ("Primary Award", rec.get("primary_award") or "No clear recommendation"),
        ("Shortlist", ", ".join(rec.get("shortlist", [])) or "—"),
        ("Reasoning", rec.get("reasoning", "")),
    ]
    for i, (label, value) in enumerate(rows, start=1):
        label_cell = ws.cell(row=i, column=1, value=label)
        label_cell.font = SUBHEADER_FONT
        value_cell = ws.cell(row=i, column=2, value=value)
        value_cell.font = DATA_FONT
        value_cell.alignment = WRAP

    caveats = report.get("caveats", [])
    if caveats:
        row = len(rows) + 2
        ws.cell(row=row, column=1, value="Caveats").font = SUBHEADER_FONT
        for c in caveats:
            row += 1
            ws.cell(row=row, column=2, value=f"• {c}").font = DATA_FONT
            ws.cell(row=row, column=2).alignment = WRAP


def _write_bidder_notes_sheet(wb, report: dict) -> None:
    ws = wb.create_sheet("Bidder Notes")
    header_row(
        ws, 1,
        ["Bidder", "Position", "Summary", "Negotiation Points", "Flags Highlighted"],
        [20, 16, 40, 50, 40],
    )
    row = 2
    for bidder, note in report.get("bidder_notes", {}).items():
        fill = GAP_FILL if note.get("position") == "excluded_data_gap" else None
        values = [
            bidder,
            note.get("position", ""),
            note.get("summary", ""),
            "\n".join(f"• {p}" for p in note.get("negotiation_points", [])),
            "\n".join(f"• {f}" for f in note.get("flags_highlighted", [])),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = DATA_FONT
            cell.alignment = WRAP
            cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
        row += 1


def _write_bid_summary_sheet(wb, result: ComparisonResult, report: dict) -> None:
    ws = wb.create_sheet("Bid Summary")
    header_row(ws, 1, ["Bidder", "Contract Total", "Rank"], [24, 20, 10])

    data_gap_bidders = set(report.get("bidder_notes", {})) & {
        b for b, n in report.get("bidder_notes", {}).items()
        if n.get("position") == "excluded_data_gap"
    }
    eligible = sorted(
        (
            (b, t) for b, t in result.bidder_grand_totals.items()
            if t is not None and b not in data_gap_bidders
        ),
        key=lambda bt: bt[1],
    )
    best_bidder = eligible[0][0] if eligible else None

    row = 2
    for bidder, rank in [(b, i + 1) for i, (b, _) in enumerate(eligible)] + [
        (b, None) for b in data_gap_bidders
    ]:
        total = result.bidder_grand_totals.get(bidder)
        cells = [
            ws.cell(row=row, column=1, value=bidder),
            ws.cell(row=row, column=2, value=total if total is not None else "—"),
            ws.cell(row=row, column=3, value=rank if rank is not None else "—"),
        ]
        for cell in cells:
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if bidder == best_bidder:
                cell.fill = BEST_PRICE_FILL
            elif bidder in data_gap_bidders:
                cell.fill = GAP_FILL
        cells[1].number_format = NUMBER_FORMAT
        row += 1


def _write_flags_sheet(wb, top_flags: dict[str, list[dict]]) -> None:
    """Same curated top-N-per-bidder flags the LLM digest saw (see
    digest.py::_top_flags) — not the full, unbounded flag list. With
    outlier detection pairwise across every bidder, the raw flag count runs
    into the thousands; dumping all of it produces a multi-megabyte,
    unreadable sheet with no relationship to what the narrative above
    actually references. This keeps the export traceable to the report."""
    ws = wb.create_sheet("Flags")
    header_row(
        ws, 1,
        ["Bidder", "Severity", "Category", "Lot", "Item No", "Description", "Detail"],
        [20, 10, 14, 16, 10, 40, 60],
    )
    row = 2
    for bidder, flags in top_flags.items():
        for flag in flags:
            values = [
                bidder, flag["severity"], flag["category"], flag["lot"],
                flag["item_no"], flag["description"], flag["detail"],
            ]
            fill = SEVERITY_FILL.get(flag["severity"])
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = DATA_FONT
                cell.alignment = WRAP
                cell.border = THIN_BORDER
                if fill:
                    cell.fill = fill
            row += 1


def build_excel_report(result: ComparisonResult, report: dict) -> BytesIO:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    digest = build_report_digest(result)

    _write_recommendation_sheet(wb, result, report)
    _write_bidder_notes_sheet(wb, report)
    _write_bid_summary_sheet(wb, result, report)
    _write_flags_sheet(wb, digest["top_flags"])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
