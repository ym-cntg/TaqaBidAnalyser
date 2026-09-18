"""Turn a parsed BoqDocument into output formats.

Two consumers: `to_dict` feeds the API and the frontend preview, and
`to_workbook` produces the normalized spreadsheet, which is the artifact
that replaces the hand-built consolidated comparison sheet described in
CLAUDE.md's manual process.
"""

from __future__ import annotations

from dataclasses import asdict
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import BoqDocument

# One flat set of columns covering both source layouts. A power sheet
# leaves Unit Rate empty; a water sheet leaves the CIF/Erection columns
# empty. Keeping one shape means downstream comparison never has to
# branch on which template a vendor happened to use.
COLUMNS = [
    ("Source Sheet", "sheet_name", 22),
    ("Section", "section", 34),
    ("Item No", "item_no", 12),
    ("Description", "description", 60),
    ("Unit", "unit", 10),
    ("Qty", "quantity", 10),
    ("CIF Unit Rate", "cif_unit_rate", 14),
    ("CIF Total", "cif_total", 14),
    ("Erection Unit Rate", "erection_unit_rate", 16),
    ("Erection Total", "erection_total", 14),
    ("Unit Rate", "unit_rate", 14),
    ("Line Total", "line_total", 16),
    ("Issue", "issue", 44),
]

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_FLAG_FILL = PatternFill("solid", fgColor="FCE4D6")
_MONEY_FORMAT = "#,##0.00"


def _line_issue(line) -> str | None:
    """Issues only apply to priced lines. A section heading or a
    subtotal has no quantity to price and no total of its own, so
    reporting "no line total" against one is noise, not a finding."""
    if line.kind != "item":
        return None
    parts = []
    if line.arithmetic_error:
        parts.append(f"Arithmetic: {line.arithmetic_error}")
    for field, text in line.non_numeric.items():
        parts.append(f"{field} reads {text!r}")
    if line.line_total is None:
        # Two very different situations, and the analyst acts on them
        # differently: a vendor who left the line blank has not quoted
        # it at all, while one who wrote "included in BOQ item 3.3" has
        # priced it somewhere else.
        parts.append("Priced elsewhere, no figure on this line" if line.non_numeric else "Unquoted")
    return "; ".join(parts) or None


def to_dict(document: BoqDocument) -> dict[str, Any]:
    reconciliation = document.reconciliation
    return {
        "source_file": document.source_file,
        "tender_ref": document.tender_ref,
        "item_count": document.item_count,
        "grand_total": document.grand_total,
        "stated_total": document.summary_total,
        "reconciliation": asdict(reconciliation) if reconciliation else None,
        "arithmetic_error_count": len(document.arithmetic_errors),
        "sheets": [
            {
                "sheet_name": sheet.sheet_name,
                "role": sheet.role,
                "layout": sheet.layout,
                "header_row": sheet.header_row,
                "item_count": sheet.item_count,
                "total": sheet.total,
                "title_lines": sheet.title_lines,
                "lines": [
                    {**asdict(line), "issue": _line_issue(line)} for line in sheet.lines
                ],
            }
            for sheet in document.sheets
        ],
        "skipped": [asdict(sheet) for sheet in document.skipped],
    }


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for index in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def to_workbook(document: BoqDocument) -> BytesIO:
    """Build the normalized BOQ workbook and return it as a stream."""
    wb = Workbook()

    boq = wb.active
    boq.title = "BOQ"
    _write_header(boq, [label for label, _, _ in COLUMNS])

    row_index = 2
    for sheet in document.boq_sheets:
        for line in sheet.lines:
            if line.kind not in ("item", "section"):
                continue
            issue = _line_issue(line)
            values = {
                "sheet_name": sheet.sheet_name,
                "section": line.section,
                "item_no": line.item_no,
                "description": line.description,
                "unit": line.unit,
                "quantity": line.quantity,
                "cif_unit_rate": line.cif_unit_rate,
                "cif_total": line.cif_total,
                "erection_unit_rate": line.erection_unit_rate,
                "erection_total": line.erection_total,
                "unit_rate": line.unit_rate,
                "line_total": line.line_total,
                "issue": issue,
            }
            for col_index, (_, key, _) in enumerate(COLUMNS, start=1):
                cell = boq.cell(row=row_index, column=col_index, value=values[key])
                if key.endswith(("_rate", "_total")) and values[key] is not None:
                    cell.number_format = _MONEY_FORMAT
            if line.kind == "section":
                for col_index in range(1, len(COLUMNS) + 1):
                    boq.cell(row=row_index, column=col_index).font = Font(bold=True)
            elif issue:
                for col_index in range(1, len(COLUMNS) + 1):
                    boq.cell(row=row_index, column=col_index).fill = _FLAG_FILL
            row_index += 1

    for col_index, (_, _, width) in enumerate(COLUMNS, start=1):
        boq.column_dimensions[get_column_letter(col_index)].width = width

    summary = wb.create_sheet("Extraction Summary")
    _write_header(summary, ["Field", "Value"])
    reconciliation = document.reconciliation
    rows: list[tuple[str, Any]] = [
        ("Source file", document.source_file),
        ("Tender reference", document.tender_ref or "not found"),
        ("Priced line items extracted", document.item_count),
        ("Extracted total", document.grand_total),
        ("Vendor stated total", document.summary_total),
    ]
    if reconciliation:
        rows += [
            ("Difference", reconciliation.difference),
            ("Difference %", round(reconciliation.percent, 4)),
            ("Reconciles", "YES" if reconciliation.agrees else "NO -- review"),
        ]
    rows.append(("Arithmetic errors flagged", len(document.arithmetic_errors)))
    for sheet in document.sheets:
        rows.append((f"Sheet [{sheet.role}] {sheet.sheet_name}", sheet.total))
    for skipped in document.skipped:
        rows.append((f"Skipped {skipped.sheet_name}", skipped.reason))

    for label, value in rows:
        summary.append([label, value])
        if isinstance(value, float):
            summary.cell(row=summary.max_row, column=2).number_format = _MONEY_FORMAT
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 30

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
