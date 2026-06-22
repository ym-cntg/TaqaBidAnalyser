from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from src.config import BOQConfig
from src.models import (
    ExtractedBOQ,
    ExtractedSummary,
    ExtractionMethod,
    LineItem,
    SectionData,
    SummaryEntry,
)


def _col_to_index(col_letter: str) -> int:
    """Convert column letter (A, B, ..., Z, AA, ...) to 1-based index."""
    result = 0
    for char in col_letter.upper():
        result = result * 26 + (ord(char) - ord("A") + 1)
    return result


INCLUDED_MARKERS = {"included", "inc", "incl", "included above", "included in"}
NOT_APPLICABLE = {"", "-", "n/a", "na", "nil", "no boq item"}


def _safe_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        cleaned = value.strip().replace(",", "")
        lower = cleaned.lower()
        if lower in NOT_APPLICABLE:
            return None
        # "Included" means price is bundled in another item — treat as explicit 0
        if lower in INCLUDED_MARKERS or lower.startswith("included"):
            return 0.0
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def _safe_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def extract_boq(
    file_path: Path,
    config: BOQConfig,
    bidder: str,
    round_name: str,
    lot: str,
) -> ExtractedBOQ:
    """Extract line items from a BOQ Excel file."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    cols = config.columns

    col_map = {
        "item_no": _col_to_index(cols.item_no),
        "description": _col_to_index(cols.description),
        "unit": _col_to_index(cols.unit),
        "qty": _col_to_index(cols.qty),
        "cif_unit_rate": _col_to_index(cols.cif_unit_rate),
        "cif_total": _col_to_index(cols.cif_total),
        "erection_unit_rate": _col_to_index(cols.erection_unit_rate),
        "erection_total": _col_to_index(cols.erection_total),
        "total_price": _col_to_index(cols.total_price),
    }

    sections: list[SectionData] = []

    for sheet_name in wb.sheetnames:
        if sheet_name.lower() in ("final summary", "manpower calc"):
            continue

        ws = wb[sheet_name]

        # Spare parts sections have fewer columns (no erection/total_price)
        effective_col_map = dict(col_map)
        if ws.max_column < _col_to_index(cols.total_price):
            # Remove columns that don't exist in this sheet
            for key in ("erection_unit_rate", "erection_total", "total_price"):
                if _col_to_index(getattr(cols, key)) > ws.max_column:
                    effective_col_map.pop(key, None)

        line_items = _extract_sheet_items(ws, effective_col_map, config.data_start_row)

        if line_items:
            sections.append(SectionData(
                name=sheet_name,
                line_items=tuple(line_items),
            ))

    wb.close()

    return ExtractedBOQ(
        bidder=bidder,
        round=round_name,
        lot=lot,
        source_file=file_path,
        extraction_method=ExtractionMethod.OPENPYXL,
        sections=tuple(sections),
    )


def _find_included_parents(
    ws,
    col_map: dict[str, int],
    data_start_row: int,
) -> set[str]:
    """Scan for 'Included in BOQ item X' patterns to find parent items.

    When a sub-item says "Included in BOQ item 3.3", all sibling sub-items
    under 3.3 are implicitly bundled into the parent's price.
    """
    parents: set[str] = set()
    price_cols = [
        col_map[f] for f in ("cif_unit_rate", "cif_total", "total_price")
        if f in col_map
    ]
    for row_num in range(data_start_row, ws.max_row + 1):
        for col_idx in price_cols:
            val = ws.cell(row=row_num, column=col_idx).value
            if val and isinstance(val, str) and "included in" in val.lower():
                m = re.search(r"item\s+(\d[\d.]*)", val, re.IGNORECASE)
                if m:
                    parents.add(m.group(1))
    return parents


def _is_sub_item_of(item_no: str, parents: set[str]) -> bool:
    """Check if item_no is a sub-item of any included parent."""
    for parent in parents:
        if item_no.startswith(parent + ".") and item_no != parent:
            return True
    return False


def _extract_sheet_items(
    ws,
    col_map: dict[str, int],
    data_start_row: int,
) -> list[LineItem]:
    """Extract line items from a single worksheet."""
    items: list[LineItem] = []
    included_parents = _find_included_parents(ws, col_map, data_start_row)

    for row_num in range(data_start_row, ws.max_row + 1):
        item_no_val = ws.cell(row=row_num, column=col_map["item_no"]).value
        desc_val = ws.cell(row=row_num, column=col_map["description"]).value

        # Skip empty rows and total/subtotal rows
        if item_no_val is None and desc_val is None:
            continue

        item_no_str = _safe_str(item_no_val)
        desc_str = _safe_str(desc_val)

        # Skip header-like rows and total rows
        if not item_no_str and not desc_str:
            continue
        if _is_total_row(desc_str, item_no_str):
            continue
        if _is_header_row(item_no_str, desc_str):
            continue

        # Determine if this is a category header (has item_no but no pricing)
        # vs an actual priced line item
        price_fields = ("cif_unit_rate", "cif_total", "erection_unit_rate",
                        "erection_total", "total_price")
        has_pricing = any(
            ws.cell(row=row_num, column=col_map[f]).value is not None
            for f in price_fields
            if f in col_map
        )
        has_qty = ws.cell(row=row_num, column=col_map["qty"]).value is not None

        # Include rows that have item numbers (even category headers for structure)
        # or rows that have pricing data
        if not item_no_str and not has_pricing:
            continue

        # Sub-items under an "Included in BOQ item X" parent with no pricing
        # are implicitly bundled — treat as 0.0 (not unquoted)
        bundled = (
            not has_pricing
            and has_qty
            and _is_sub_item_of(item_no_str, included_parents)
        )

        items.append(LineItem(
            item_no=item_no_str,
            description=desc_str,
            unit=_safe_str(ws.cell(row=row_num, column=col_map["unit"]).value) or None,
            qty=_safe_float(ws.cell(row=row_num, column=col_map["qty"]).value),
            cif_unit_rate=0.0 if bundled else (_safe_float(
                ws.cell(row=row_num, column=col_map["cif_unit_rate"]).value
            ) if "cif_unit_rate" in col_map else None),
            cif_total=0.0 if bundled else (_safe_float(
                ws.cell(row=row_num, column=col_map["cif_total"]).value
            ) if "cif_total" in col_map else None),
            erection_unit_rate=0.0 if bundled else (_safe_float(
                ws.cell(row=row_num, column=col_map["erection_unit_rate"]).value
            ) if "erection_unit_rate" in col_map else None),
            erection_total=0.0 if bundled else (_safe_float(
                ws.cell(row=row_num, column=col_map["erection_total"]).value
            ) if "erection_total" in col_map else None),
            total_price=0.0 if bundled else (_safe_float(
                ws.cell(row=row_num, column=col_map["total_price"]).value
            ) if "total_price" in col_map else None),
        ))

    return items


def _is_total_row(desc: str, item_no: str = "") -> bool:
    lower = desc.lower().strip()
    lower_item = item_no.lower().strip()

    # Check description for total keywords
    if any(kw in lower for kw in (
        "total price", "grand total", "sub total", "subtotal",
        "total lot", "total contract", "u.a.e. dirhams",
    )):
        return True

    # Check item_no field for subtotal patterns
    # e.g., "Item-1 Sub Total ...", "GRAND TOTAL of Items-1,2,3"
    if any(kw in lower_item for kw in (
        "sub total", "grand total", "total",
    )):
        return True

    # Recap rows: "Item - 1", "Item - 2", "Item - 3" at end of sheet
    if re.match(r"^item\s*-?\s*\d+$", lower_item):
        return True

    return False


def _is_header_row(item_no: str, desc: str) -> bool:
    lower_item = item_no.lower().strip()
    lower_desc = desc.lower().strip()
    return (
        lower_item in ("item", "item no", "item no.")
        or lower_desc in ("description",)
        or "technical specifications" in lower_desc
    )


def extract_summary(
    file_path: Path,
    config: BOQConfig,
    bidder: str,
    round_name: str,
) -> ExtractedSummary:
    """Extract the summary file (combined lot totals)."""
    wb = openpyxl.load_workbook(file_path, data_only=True)

    # Find the Final Summary sheet
    summary_sheet = None
    for name in wb.sheetnames:
        if "summary" in name.lower():
            summary_sheet = wb[name]
            break

    if summary_sheet is None:
        summary_sheet = wb[wb.sheetnames[0]]

    lots_data: dict[str, tuple[SummaryEntry, ...]] = {}
    lot_totals: dict[str, float] = {}

    for lot_id, (cif_col, erect_col, total_col) in config.summary_lot_columns.items():
        cif_idx = _col_to_index(cif_col)
        erect_idx = _col_to_index(erect_col)
        total_idx = _col_to_index(total_col)

        entries: list[SummaryEntry] = []
        for row_num in range(config.summary_data_start_row, summary_sheet.max_row + 1):
            item_val = summary_sheet.cell(row=row_num, column=1).value
            desc_val = summary_sheet.cell(row=row_num, column=2).value

            if item_val is None and desc_val is None:
                continue

            desc_str = _safe_str(desc_val)

            # Capture lot total
            if "total price lot" in desc_str.lower() or "total lot" in desc_str.lower():
                total_val = _safe_float(
                    summary_sheet.cell(row=row_num, column=total_idx).value
                )
                if total_val is not None:
                    lot_totals[lot_id] = total_val
                continue

            if item_val is None:
                continue

            entries.append(SummaryEntry(
                item_no=_safe_str(item_val),
                description=desc_str,
                cif_total=_safe_float(
                    summary_sheet.cell(row=row_num, column=cif_idx).value
                ),
                erection_total=_safe_float(
                    summary_sheet.cell(row=row_num, column=erect_idx).value
                ),
                total_price=_safe_float(
                    summary_sheet.cell(row=row_num, column=total_idx).value
                ),
            ))

        if entries:
            lots_data[lot_id] = tuple(entries)

    # Contract total
    contract_total = None
    for row_num in range(config.summary_data_start_row, summary_sheet.max_row + 1):
        desc_val = _safe_str(
            summary_sheet.cell(row=row_num, column=2).value
        )
        if "total contract" in desc_val.lower():
            # Contract total is usually in a merged area — check multiple columns
            for col in (3, 5, 6):
                val = _safe_float(summary_sheet.cell(row=row_num, column=col).value)
                if val is not None:
                    contract_total = val
                    break
            break

    wb.close()

    return ExtractedSummary(
        bidder=bidder,
        round=round_name,
        source_file=file_path,
        extraction_method=ExtractionMethod.OPENPYXL,
        lots=lots_data,
        lot_totals=lot_totals,
        contract_total=contract_total,
    )
