from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.analysis.comparator import LotComparison, build_tender_summary, compare_lot
from src.analysis.round_tracker import TenderRoundSummary
from src.models import Finding, FindingCategory, Severity, TenderData

# Style constants
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
SUBHEADER_FONT = Font(name="Calibri", bold=True, size=10)
DATA_FONT = Font(name="Calibri", size=10)
NUMBER_FORMAT = "#,##0"
RATE_FORMAT = "#,##0.00"
OUTLIER_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
UNQUOTED_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
BEST_PRICE_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def generate_excel_report(
    tender: TenderData,
    round_name: str,
    round_summary: TenderRoundSummary,
    output_path: Path,
) -> Path:
    """Generate the consolidated comparison Excel workbook."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Summary sheet first
    _write_summary_sheet(wb, tender, round_name, round_summary)

    # One sheet per lot with full line-item comparison
    for lot in tender.lots:
        comparison = compare_lot(tender, lot, round_name)
        _write_lot_sheet(wb, comparison, tender.findings)

    # Findings sheet
    _write_findings_sheet(wb, tender.findings)

    wb.save(output_path)
    return output_path


def _write_summary_sheet(
    wb: openpyxl.Workbook,
    tender: TenderData,
    round_name: str,
    round_summary: TenderRoundSummary,
) -> None:
    ws = wb.create_sheet("Summary")

    # Title
    ws.merge_cells("A1:H1")
    ws["A1"] = f"Bid Comparison Summary — {tender.tender_id} ({round_name})"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14)

    # Lot totals table
    row = 3
    bidders = tender.bidders
    lots = tender.lots

    # Headers
    ws.cell(row=row, column=1, value="Bidder").font = HEADER_FONT
    ws.cell(row=row, column=1).fill = HEADER_FILL
    for j, lot in enumerate(lots):
        col = 2 + j
        ws.cell(row=row, column=col, value=lot).font = HEADER_FONT
        ws.cell(row=row, column=col).fill = HEADER_FILL
    total_col = 2 + len(lots)
    ws.cell(row=row, column=total_col, value="Contract Total").font = HEADER_FONT
    ws.cell(row=row, column=total_col).fill = HEADER_FILL
    rank_col = total_col + 1
    ws.cell(row=row, column=rank_col, value="Rank").font = HEADER_FONT
    ws.cell(row=row, column=rank_col).fill = HEADER_FILL

    summary = build_tender_summary(tender, round_name)

    # Calculate totals for ranking
    bidder_totals: dict[str, float] = {}
    for bidder in bidders:
        lot_totals = summary.get(bidder, {})
        grand = sum(v for v in lot_totals.values() if v is not None)
        bidder_totals[bidder] = grand

    ranked = sorted(bidder_totals.items(), key=lambda x: x[1] if x[1] > 0 else float("inf"))
    rank_map = {bidder: i + 1 for i, (bidder, _) in enumerate(ranked)}

    # Data rows
    for i, bidder in enumerate(bidders):
        row = 4 + i
        ws.cell(row=row, column=1, value=bidder).font = DATA_FONT
        ws.cell(row=row, column=1).border = THIN_BORDER

        for j, lot in enumerate(lots):
            col = 2 + j
            val = summary.get(bidder, {}).get(lot)
            cell = ws.cell(row=row, column=col, value=val)
            cell.number_format = NUMBER_FORMAT
            cell.font = DATA_FONT
            cell.border = THIN_BORDER

        # Contract total
        total = bidder_totals.get(bidder, 0)
        cell = ws.cell(row=row, column=total_col, value=total if total > 0 else None)
        cell.number_format = NUMBER_FORMAT
        cell.font = Font(name="Calibri", bold=True, size=10)
        cell.border = THIN_BORDER

        # Highlight lowest
        if rank_map.get(bidder) == 1 and total > 0:
            cell.fill = BEST_PRICE_FILL

        # Rank
        cell = ws.cell(row=row, column=rank_col, value=rank_map.get(bidder))
        cell.font = DATA_FONT
        cell.border = THIN_BORDER

    # Round movement section
    row = 4 + len(bidders) + 2
    ws.cell(row=row, column=1, value="Price Movement Across Rounds").font = Font(
        name="Calibri", bold=True, size=12
    )
    row += 1

    headers = ["Bidder", "Lot"] + list(tender.rounds) + ["Total Change %"]
    for j, h in enumerate(headers):
        cell = ws.cell(row=row, column=j + 1, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for movement in round_summary.movements:
        row += 1
        ws.cell(row=row, column=1, value=movement.bidder).font = DATA_FONT
        ws.cell(row=row, column=2, value=movement.lot).font = DATA_FONT

        for j, rnd in enumerate(tender.rounds):
            val = movement.round_totals.get(rnd)
            cell = ws.cell(row=row, column=3 + j, value=val)
            cell.number_format = NUMBER_FORMAT
            cell.font = DATA_FONT

        pct = movement.total_change_pct
        if pct is not None:
            cell = ws.cell(row=row, column=3 + len(tender.rounds), value=f"{pct:+.1f}%")
            cell.font = DATA_FONT

    # Auto-width
    _auto_width(ws)


def _write_lot_sheet(
    wb: openpyxl.Workbook,
    comparison: LotComparison,
    findings: list[Finding],
) -> None:
    sheet_name = comparison.lot[:31]  # Excel max 31 chars
    ws = wb.create_sheet(sheet_name)

    bidders = comparison.bidders
    finding_set = _build_finding_set(findings, comparison.lot, comparison.round)

    row = 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5 + len(bidders) * 3)
    ws["A1"] = f"{comparison.lot} — {comparison.round} — Line Item Comparison"
    ws["A1"].font = Font(name="Calibri", bold=True, size=12)

    for section in comparison.sections:
        row += 2
        ws.cell(row=row, column=1, value=section.section_name).font = Font(
            name="Calibri", bold=True, size=11, color="1F4E79"
        )

        # Column headers
        row += 1
        base_headers = ["Item No.", "Description", "Unit", "Qty"]
        for j, h in enumerate(base_headers):
            cell = ws.cell(row=row, column=j + 1, value=h)
            cell.font = SUBHEADER_FONT
            cell.fill = SUBHEADER_FILL
            cell.border = THIN_BORDER

        col = len(base_headers) + 1
        for bidder in bidders:
            for sub_h in ("CIF Total", "Erection Total", "Total Price"):
                cell = ws.cell(row=row, column=col, value=f"{bidder}\n{sub_h}")
                cell.font = SUBHEADER_FONT
                cell.fill = SUBHEADER_FILL
                cell.border = THIN_BORDER
                cell.alignment = Alignment(wrap_text=True, horizontal="center")
                col += 1

        # Median column
        cell = ws.cell(row=row, column=col, value="Median\nTotal")
        cell.font = SUBHEADER_FONT
        cell.fill = SUBHEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(wrap_text=True, horizontal="center")

        # Data rows
        for item in section.items:
            row += 1
            ws.cell(row=row, column=1, value=item.item_no).font = DATA_FONT
            ws.cell(row=row, column=1).border = THIN_BORDER

            desc_cell = ws.cell(row=row, column=2, value=item.description)
            desc_cell.font = DATA_FONT
            desc_cell.border = THIN_BORDER
            desc_cell.alignment = Alignment(wrap_text=True)

            ws.cell(row=row, column=3, value=item.unit).font = DATA_FONT
            ws.cell(row=row, column=3).border = THIN_BORDER

            qty_cell = ws.cell(row=row, column=4, value=item.qty)
            qty_cell.number_format = NUMBER_FORMAT
            qty_cell.font = DATA_FONT
            qty_cell.border = THIN_BORDER

            col = 5
            min_price = item.min_total
            for bidder in bidders:
                cif = item.bidder_cif_total.get(bidder)
                erect = item.bidder_erection_total.get(bidder)
                total = item.bidder_total_price.get(bidder)

                for val in (cif, erect, total):
                    cell = ws.cell(row=row, column=col, value=val)
                    cell.number_format = NUMBER_FORMAT
                    cell.font = DATA_FONT
                    cell.border = THIN_BORDER
                    col += 1

                # Highlight
                total_cell = ws.cell(row=row, column=col - 1)
                finding_key = (bidder, item.item_no)
                if finding_key in finding_set:
                    for cat in finding_set[finding_key]:
                        if cat == FindingCategory.OUTLIER:
                            total_cell.fill = OUTLIER_FILL
                        elif cat == FindingCategory.UNQUOTED:
                            total_cell.fill = UNQUOTED_FILL

                # Best price highlight
                if total is not None and min_price is not None and total == min_price:
                    total_cell.fill = BEST_PRICE_FILL

            # Median
            cell = ws.cell(row=row, column=col, value=item.median_total)
            cell.number_format = NUMBER_FORMAT
            cell.font = Font(name="Calibri", italic=True, size=10)
            cell.border = THIN_BORDER

    # Lot totals row
    row += 2
    ws.cell(row=row, column=1, value="LOT TOTAL").font = Font(
        name="Calibri", bold=True, size=11
    )
    col = 5
    for bidder in bidders:
        col += 2  # skip CIF and Erection columns
        total = comparison.bidder_lot_totals.get(bidder)
        cell = ws.cell(row=row, column=col, value=total)
        cell.number_format = NUMBER_FORMAT
        cell.font = Font(name="Calibri", bold=True, size=11)
        cell.border = THIN_BORDER
        col += 1

    _auto_width(ws)


def _write_findings_sheet(
    wb: openpyxl.Workbook,
    findings: list[Finding],
) -> None:
    ws = wb.create_sheet("Findings")

    headers = ["Category", "Severity", "Bidder", "Round", "Lot", "Item No.", "Finding"]
    for j, h in enumerate(headers):
        cell = ws.cell(row=1, column=j + 1, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    severity_order = {"high": 0, "medium": 1, "low": 2}
    sorted_findings = sorted(
        findings,
        key=lambda f: (severity_order.get(f.severity.value, 99), f.category.value),
    )

    for i, finding in enumerate(sorted_findings):
        row = i + 2
        ws.cell(row=row, column=1, value=finding.category.value).font = DATA_FONT
        sev_cell = ws.cell(row=row, column=2, value=finding.severity.value)
        sev_cell.font = DATA_FONT
        if finding.severity == Severity.HIGH:
            sev_cell.fill = OUTLIER_FILL
        ws.cell(row=row, column=3, value=finding.bidder).font = DATA_FONT
        ws.cell(row=row, column=4, value=finding.round).font = DATA_FONT
        ws.cell(row=row, column=5, value=finding.lot).font = DATA_FONT
        ws.cell(row=row, column=6, value=finding.item_no or "—").font = DATA_FONT
        ws.cell(row=row, column=7, value=finding.message).font = DATA_FONT

    _auto_width(ws)


def _build_finding_set(
    findings: list[Finding],
    lot: str,
    round_name: str,
) -> dict[tuple[str, str | None], list[FindingCategory]]:
    """Build lookup: (bidder, item_no) -> list of finding categories."""
    result: dict[tuple[str, str | None], list[FindingCategory]] = {}
    for f in findings:
        if f.lot != lot or f.round != round_name:
            continue
        key = (f.bidder, f.item_no)
        if key not in result:
            result[key] = []
        result[key].append(f.category)
    return result


def _auto_width(ws) -> None:
    for col_cells in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                val_str = str(cell.value)
                lines = val_str.split("\n")
                max_line = max(len(line) for line in lines)
                max_length = max(max_length, max_line)
        adjusted = min(max_length + 2, 40)
        ws.column_dimensions[col_letter].width = adjusted
