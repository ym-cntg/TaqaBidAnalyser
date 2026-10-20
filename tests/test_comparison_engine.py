"""The comparison, round-tracking and negotiation-report engines.

Two things are guarded here. That the Maximo path is unchanged by the
source injection added for uploads, and that the report prices at the
latest round rather than the original quote.
"""

from __future__ import annotations

from backend.api.comparison import build_comparison
from backend.api.negotiation_report import build_negotiation_report, _build_excel
from backend.api.rounds import build_round_trend
from tests.conftest import make_uploaded

ORIGINAL = "original"


class TestMaximoPathUnchanged:
    """The source argument is additive. Without one, every fetch must
    still hit Maximo exactly as it did before uploads existed."""

    def test_maximo_fetches_are_still_called(self, maximo):
        build_comparison("D-111808")
        assert "roster" in maximo and "lines" in maximo

    def test_comparison_defaults_to_the_original_round(self, maximo):
        assert build_comparison("D-111808")["round"] == ORIGINAL

    def test_original_totals_intact(self, maximo):
        res = build_comparison("D-111808")
        assert {v["vendor"]: v["contract_total"] for v in res["vendors"]} == {
            "A": 200.0,
            "B": 260.0,
        }

    def test_a_later_round_reprices_from_discount_history(self, maximo):
        res = build_comparison("D-111808", round_label="1.0")
        assert {v["vendor"]: v["contract_total"] for v in res["vendors"]} == {
            "A": 150.0,
            "B": 120.0,
        }


class TestUploadedSourceDrivesTheSameEngines:
    def test_comparison_from_an_uploaded_boq(self, no_overlays):
        src = make_uploaded({ORIGINAL: {"ALPHA": 100.0, "BETA": 150.0}})
        res = build_comparison("p", source=src)
        assert sorted(v["vendor"] for v in res["vendors"]) == ["ALPHA", "BETA"]
        # 2 lines x qty 10 x unit cost
        assert {v["vendor"]: v["contract_total"] for v in res["vendors"]} == {
            "ALPHA": 2000.0,
            "BETA": 3000.0,
        }

    def test_rounds_are_reconstructed_from_uploads(self, no_overlays):
        src = make_uploaded({ORIGINAL: {"ALPHA": 100.0}, "1.0": {"ALPHA": 90.0}})
        res = build_comparison("p", source=src)
        assert res["rounds_present"] == [ORIGINAL, "1.0"]
        later = build_comparison("p", round_label="1.0", source=src)
        assert later["vendors"][0]["contract_total"] == 1800.0

    def test_round_trend_sees_both_rounds(self, no_overlays):
        src = make_uploaded({ORIGINAL: {"ALPHA": 100.0, "BETA": 150.0},
                             "1.0": {"ALPHA": 90.0, "BETA": 95.0}})
        trend = build_round_trend("p", source=src)
        assert trend["rounds_present"] == [ORIGINAL, "1.0"]
        assert len(trend["vendors"]) == 2

    def test_cheapest_flag_computed_on_uploaded_data(self, no_overlays):
        src = make_uploaded({ORIGINAL: {"ALPHA": 100.0, "BETA": 150.0}})
        res = build_comparison("p", source=src)
        assert res["lines"][0]["prices"]["ALPHA"]["is_lowest"] is True
        assert res["lines"][0]["prices"]["BETA"]["is_lowest"] is False


class TestReportUsesTheLatestRound:
    """Regression. The report called build_comparison() with no round,
    which silently defaults to the original quote, so every total and
    ranking ignored every discount since. On a tender where a vendor
    discounts hard and overtakes, it named the wrong vendor as lowest."""

    def _overtaking(self):
        # ALPHA opens cheapest; BETA discounts hard and overtakes by the
        # final round.
        return make_uploaded({
            ORIGINAL: {"ALPHA": 100.0, "BETA": 150.0},   # 2000 vs 3000
            "1.0":    {"ALPHA":  95.0, "BETA":  90.0},   # 1900 vs 1800
            "2.0":    {"ALPHA":  94.0, "BETA":  70.0},   # 1880 vs 1400
        })

    def test_report_states_which_round_it_used(self, no_overlays):
        rep = build_negotiation_report("p", source=self._overtaking())
        assert rep["round"] == "2.0"
        assert rep["rounds_present"] == [ORIGINAL, "1.0", "2.0"]

    def test_totals_are_the_final_round(self, no_overlays):
        rep = build_negotiation_report("p", source=self._overtaking())
        assert {v["vendor"]: v["contract_total"] for v in rep["vendors"]} == {
            "ALPHA": 1880.0,
            "BETA": 1400.0,
        }

    def test_ranking_follows_the_final_round(self, no_overlays):
        rep = build_negotiation_report("p", source=self._overtaking())
        ranks = {v["vendor"]: v["rank"] for v in rep["vendors"]}
        assert ranks["BETA"] == 1 and ranks["ALPHA"] == 2

    def test_pct_above_lowest_uses_final_round_totals(self, no_overlays):
        rep = build_negotiation_report("p", source=self._overtaking())
        alpha = next(v for v in rep["vendors"] if v["vendor"] == "ALPHA")
        assert abs(alpha["pct_above_lowest"] - 34.29) < 0.1

    def test_single_round_tender_is_unchanged(self, no_overlays):
        src = make_uploaded({ORIGINAL: {"ALPHA": 100.0, "BETA": 150.0}})
        rep = build_negotiation_report("p", source=src)
        assert rep["round"] == ORIGINAL
        assert rep["has_round_data"] is False
        assert {v["vendor"]: v["contract_total"] for v in rep["vendors"]} == {
            "ALPHA": 2000.0,
            "BETA": 3000.0,
        }

    def test_report_follows_the_latest_round_on_the_maximo_path_too(self, maximo):
        rep = build_negotiation_report("D-111808")
        assert rep["round"] == "1.0"
        assert {v["vendor"]: v["contract_total"] for v in rep["vendors"]} == {
            "A": 150.0,
            "B": 120.0,
        }

    def test_excel_export_states_the_round(self, no_overlays):
        """The workbook gets mailed around detached from the page that
        produced it, so the round has to be on the sheet."""
        import io
        import openpyxl

        rep = build_negotiation_report("p", source=self._overtaking())
        ws = openpyxl.load_workbook(io.BytesIO(_build_excel(rep).getvalue()))["Summary"]
        assert "Round 2" in (ws.cell(row=1, column=1).value or "")
        assert ws.cell(row=3, column=1).value == "Vendor"
