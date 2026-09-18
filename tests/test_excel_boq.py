"""Verification for the Excel BOQ extractor.

Run: python -m tests.test_excel_boq

Deliberately a plain script with asserts rather than a pytest suite, to
match the "no committed test framework" state of this repo, and it runs
against the real vendor files in data/ rather than synthetic fixtures,
because every hard part of this parser (merged multi-row headers,
subtotal rows, cross-referenced prices) only exists in real files.
"""

from __future__ import annotations

import glob
import sys
from io import BytesIO

from backend.excel_boq.generator import to_dict, to_workbook
from backend.excel_boq.parser import parse_workbook

POWER = "data/power/AGPOWER/original/excel/D-111808 BOQ-Lot 1-SHBPRY.xlsx"
WATER = "data/water/A-20669/003108-PURE WATER TECHNOLOGY. LLC/PWT-BestAndFinalOffer-A-20669pricedBOQ.xlsx"
# A real file whose own summary sheet contradicts its priced detail.
INCONSISTENT = "data/power/Ray/round1/excel/Lot-1-SHBPRY-R1.xlsx"

failures: list[str] = []


def check(label: str, condition: bool, detail: str = "") -> None:
    if condition:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label} {detail}")
        failures.append(label)


def test_power_layout() -> None:
    print("\nADDC power template")
    doc = parse_workbook(POWER, "lot1.xlsx")

    check("tender reference read from the title block", doc.tender_ref == "D-111808", doc.tender_ref or "")

    boq = doc.boq_sheets[0]
    check("CIF and Erection mapped as separate columns", boq.layout == "power")
    check(
        "two-row merged header resolved",
        boq.column_map.cif_unit_rate == 4 and boq.column_map.erection_unit_rate == 6,
        str(boq.column_map),
    )

    # The published contract total for this file, from its own summary.
    check(
        "extracted detail equals the vendor's stated total",
        abs(doc.grand_total - 66_200_810) < 1,
        f"{doc.grand_total:,.2f}",
    )
    check("reconciliation agrees", doc.reconciliation is not None and doc.reconciliation.agrees)

    # Sheet-level breakdown must match the Final Summary line by line,
    # which is what proves subtotal rows are not being double-counted.
    expected = {
        "Items 1,2,3-Construction Works": 58_256_990,
        "Items 4-Load Diversion Works": 2_464_540,
        "Items 5-8 Dismantl & Modif work": 3_876_840,
        "Items 9- Spare Parts (Optional)": 1_602_440,
    }
    for sheet in doc.boq_sheets:
        check(
            f"sheet total matches summary: {sheet.sheet_name[:30]}",
            abs(sheet.total - expected[sheet.sheet_name]) < 1,
            f"{sheet.total:,.2f} != {expected[sheet.sheet_name]:,.2f}",
        )

    first = next(l for l in boq.lines if l.kind == "item")
    check(
        "first priced line read correctly",
        first.item_no == "1.1.1"
        and first.quantity == 3
        and first.cif_unit_rate == 448_080
        and first.erection_total == 268_860
        and first.line_total == 1_613_100,
        str(first),
    )
    check("section heading carried down onto its children", first.section == "1.1 33 kV GIS Switchgear", str(first.section))
    check("summary sheet excluded from the document total", len(doc.summary_sheets) == 1)


def test_water_layout() -> None:
    print("\nWater bill template")
    doc = parse_workbook(WATER, WATER.split("/")[-1])

    check("tender reference falls back to the filename", doc.tender_ref == "A-20669", str(doc.tender_ref))
    check("single-rate layout detected", doc.boq_sheets[0].layout == "water")
    check("every bill sheet parsed", len(doc.boq_sheets) == 22, str(len(doc.boq_sheets)))
    check(
        "nested summaries do not double-count",
        doc.summary_total is not None and abs(doc.summary_total - 61_247_740) < 1,
        f"{doc.summary_total:,.2f}",
    )
    check(
        "detail reconciles with the project summary",
        doc.reconciliation is not None and doc.reconciliation.agrees,
        f"{doc.reconciliation.percent:.3f}%" if doc.reconciliation else "",
    )


def test_cross_references_not_coerced() -> None:
    print("\nCross-referenced prices")
    doc = parse_workbook(INCONSISTENT, "r1.xlsx")
    sheet = next(s for s in doc.sheets if s.sheet_name.startswith("Items 1,2,3"))
    line = next(l for l in sheet.lines if l.item_no == "3.3.1")

    # "Included in BOQ item 3.3" must never become the number 3.3.
    check("text in a money column is not coerced to a number", line.cif_unit_rate is None, str(line.cif_unit_rate))
    check(
        "the cross-reference text is preserved",
        "Included in BOQ item 3.3" in line.non_numeric.get("cif_unit_rate", ""),
        str(line.non_numeric),
    )
    issue = to_dict(doc)
    flagged = [l for s in issue["sheets"] for l in s["lines"] if l["issue"]]
    check("cross-referenced lines are flagged", len(flagged) > 0, str(len(flagged)))
    check(
        "issues are never raised against section headings",
        all(l["kind"] == "item" for s in issue["sheets"] for l in s["lines"] if l["issue"]),
    )


def test_inconsistent_file_is_caught() -> None:
    print("\nVendor self-inconsistency")
    doc = parse_workbook(INCONSISTENT, "r1.xlsx")
    check("mismatch is detected", doc.reconciliation is not None and not doc.reconciliation.agrees)
    check(
        "mismatch is quantified",
        doc.reconciliation is not None and abs(doc.reconciliation.percent - 81.05) < 0.1,
        f"{doc.reconciliation.percent:.2f}%" if doc.reconciliation else "",
    )


def test_export_roundtrip() -> None:
    print("\nWorkbook export")
    import openpyxl

    doc = parse_workbook(POWER, "lot1.xlsx")
    stream = to_workbook(doc)
    wb = openpyxl.load_workbook(BytesIO(stream.getvalue()))
    check("export has both sheets", wb.sheetnames == ["BOQ", "Extraction Summary"], str(wb.sheetnames))

    ws = wb["BOQ"]
    check("export has one row per item and section", ws.max_row > doc.item_count, str(ws.max_row))
    headers = [c.value for c in ws[1]]
    check("export header is the normalized column set", headers[0] == "Source Sheet" and "Line Total" in headers)

    total_col = headers.index("Line Total") + 1
    exported = sum(
        ws.cell(row=r, column=total_col).value or 0
        for r in range(2, ws.max_row + 1)
    )
    check(
        "exported totals equal the parsed totals",
        abs(exported - doc.grand_total) < 1,
        f"{exported:,.2f} != {doc.grand_total:,.2f}",
    )


def test_corpus_does_not_crash() -> None:
    print("\nFull sample corpus")
    files = sorted(glob.glob("data/**/*.xlsx", recursive=True))
    parsed = reconciled = checked = 0
    for path in files:
        try:
            doc = parse_workbook(path, path.split("/")[-1])
        except Exception as exc:
            check(f"parse {path.split('/')[-1]}", False, f"{type(exc).__name__}: {exc}")
            continue
        parsed += 1
        if doc.reconciliation:
            checked += 1
            reconciled += int(doc.reconciliation.agrees)

    check(f"all {len(files)} sample files parse", parsed == len(files), f"{parsed}/{len(files)}")
    # The two known exceptions are one real vendor file, in two rounds,
    # whose own summary sheet is internally inconsistent.
    check(
        "all but the two known-inconsistent files reconcile",
        reconciled == checked - 2,
        f"{reconciled}/{checked}",
    )


def main() -> int:
    for test in (
        test_power_layout,
        test_water_layout,
        test_cross_references_not_coerced,
        test_inconsistent_file_is_caught,
        test_export_roundtrip,
        test_corpus_does_not_crash,
    ):
        test()

    print(f"\n{'FAILED: ' + ', '.join(failures) if failures else 'All checks passed.'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
