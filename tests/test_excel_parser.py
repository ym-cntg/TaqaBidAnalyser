"""Excel BOQ extraction, against real vendor spreadsheets.

Synthetic fixtures were deliberately not used: every hard part of this
parser (merged multi-row headers, subtotal rows, cross-referenced
prices) only exists in files a real bidder produced.
"""

from __future__ import annotations

import glob
from io import BytesIO

import openpyxl
import pytest

from backend.excel_boq.generator import to_dict, to_workbook
from backend.excel_boq.parser import parse_workbook
from tests.conftest import (
    DATA_DIR,
    POWER_V1_ORIGINAL,
    POWER_V2_ROUND1,
    WATER_ORIGINAL,
    needs_full_corpus,
)


@pytest.fixture(scope="module")
def power():
    return parse_workbook(POWER_V1_ORIGINAL, POWER_V1_ORIGINAL.name)


@pytest.fixture(scope="module")
def water():
    return parse_workbook(WATER_ORIGINAL, WATER_ORIGINAL.name)


@pytest.fixture(scope="module")
def inconsistent():
    return parse_workbook(POWER_V2_ROUND1, POWER_V2_ROUND1.name)


class TestPowerTemplate:
    def test_tender_reference_from_title_block(self, power):
        assert power.tender_ref == "D-111808"

    def test_cif_and_erection_mapped_separately(self, power):
        """The two-row merged header puts the group labels one row above
        the sub-labels. Reading only the scored row collapses two
        different Unit Rate columns into one and loses the split."""
        boq = power.boq_sheets[0]
        assert boq.layout == "power"
        assert boq.column_map.cif_unit_rate == 4
        assert boq.column_map.erection_unit_rate == 6

    def test_total_matches_the_vendors_own_summary(self, power):
        assert abs(power.grand_total - 66_200_810) < 1
        assert power.reconciliation is not None
        assert power.reconciliation.agrees

    @pytest.mark.parametrize(
        "sheet_name,expected",
        [
            ("Items 1,2,3-Construction Works", 58_256_990),
            ("Items 4-Load Diversion Works", 2_464_540),
            ("Items 5-8 Dismantl & Modif work", 3_876_840),
            ("Items 9- Spare Parts (Optional)", 1_602_440),
        ],
    )
    def test_each_sheet_matches_the_summary_breakdown(self, power, sheet_name, expected):
        """Proves subtotal rows are not double counted. Counting them
        inflated this file from 66.2M to 233M."""
        sheet = next(s for s in power.boq_sheets if s.sheet_name == sheet_name)
        assert abs(sheet.total - expected) < 1

    def test_first_priced_line_read_correctly(self, power):
        line = next(l for l in power.boq_sheets[0].lines if l.kind == "item")
        assert line.item_no == "1.1.1"
        assert line.quantity == 3
        assert line.cif_unit_rate == 448_080
        assert line.erection_total == 268_860
        assert line.line_total == 1_613_100

    def test_section_heading_carried_onto_children(self, power):
        line = next(l for l in power.boq_sheets[0].lines if l.kind == "item")
        assert line.section == "1.1 33 kV GIS Switchgear"

    def test_summary_sheet_excluded_from_the_total(self, power):
        assert len(power.summary_sheets) == 1


class TestWaterTemplate:
    def test_tender_reference_falls_back_to_filename(self, water):
        assert water.tender_ref == "A-20669"

    def test_single_rate_layout(self, water):
        assert water.boq_sheets[0].layout == "water"

    def test_every_bill_sheet_parsed(self, water):
        assert len(water.boq_sheets) == 22

    def test_nested_summaries_do_not_double_count(self, water):
        """This workbook has a project summary plus one per area. Summing
        them all would double count, so the top-level one is used."""
        assert water.summary_total is not None
        assert abs(water.summary_total - 61_247_740) < 1

    def test_detail_reconciles_with_the_project_summary(self, water):
        assert water.reconciliation is not None
        assert water.reconciliation.agrees


class TestCrossReferencedPrices:
    def test_text_in_a_money_column_is_not_coerced(self, inconsistent):
        """"Included in BOQ item 3.3" must never become the number 3.3.
        Stripping letters indiscriminately invented a price no vendor
        quoted."""
        sheet = next(s for s in inconsistent.sheets if s.sheet_name.startswith("Items 1,2,3"))
        line = next(l for l in sheet.lines if l.item_no == "3.3.1")
        assert line.cif_unit_rate is None
        assert "Included in BOQ item 3.3" in line.non_numeric.get("cif_unit_rate", "")

    def test_cross_referenced_lines_are_flagged(self, inconsistent):
        payload = to_dict(inconsistent)
        flagged = [l for s in payload["sheets"] for l in s["lines"] if l["issue"]]
        assert flagged

    def test_issues_never_raised_against_section_headings(self, inconsistent):
        payload = to_dict(inconsistent)
        assert all(
            l["kind"] == "item"
            for s in payload["sheets"]
            for l in s["lines"]
            if l["issue"]
        )


class TestVendorSelfInconsistency:
    def test_mismatch_detected_and_quantified(self, inconsistent):
        """A real file whose own summary shows zero for items 3 to 8
        while the detail carries 22.5M of priced work."""
        rec = inconsistent.reconciliation
        assert rec is not None
        assert not rec.agrees
        assert abs(rec.percent - 81.05) < 0.1


class TestWorkbookExport:
    def test_export_structure_and_totals(self, power):
        wb = openpyxl.load_workbook(BytesIO(to_workbook(power).getvalue()))
        assert wb.sheetnames == ["BOQ", "Extraction Summary"]
        ws = wb["BOQ"]
        headers = [c.value for c in ws[1]]
        assert headers[0] == "Source Sheet"
        assert "Line Total" in headers

        total_col = headers.index("Line Total") + 1
        exported = sum(
            ws.cell(row=r, column=total_col).value or 0 for r in range(2, ws.max_row + 1)
        )
        assert abs(exported - power.grand_total) < 1


@needs_full_corpus
class TestFullCorpus:
    """Runs only where data/ is present, which is a developer machine.
    CI covers the committed fixtures instead."""

    def test_every_sample_file_parses(self):
        files = sorted(glob.glob(str(DATA_DIR / "**" / "*.xlsx"), recursive=True))
        assert files
        for path in files:
            parse_workbook(path, path.split("/")[-1])

    def test_all_but_the_known_inconsistent_files_reconcile(self):
        files = sorted(glob.glob(str(DATA_DIR / "**" / "*.xlsx"), recursive=True))
        checked = reconciled = 0
        for path in files:
            doc = parse_workbook(path, path.split("/")[-1])
            if doc.reconciliation:
                checked += 1
                reconciled += int(doc.reconciliation.agrees)
        # One real vendor file across two rounds is genuinely wrong.
        assert reconciled == checked - 2
